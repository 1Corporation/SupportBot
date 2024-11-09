import io
import datetime

import openpyxl
from openpyxl.styles import Alignment, PatternFill
from openpyxl.styles.colors import Color, COLOR_INDEX
from openpyxl.utils.cell import get_column_letter
from aiogram.filters import Command
from aiogram.types import Message, BufferedInputFile

from bot.dispatcher import dp
from utils import DatabaseConnection


@dp.message(Command("month_excel"))
async def month_excel_handler(message: Message):
    await month_excel(message)


async def month_excel(message: Message) -> None:
    """
    Генерирует excel выгрузку за месяц и отправляет в чат
    :param message: aigram message Object
    :return: None
    """
    # Создаем таблицу
    wb = openpyxl.Workbook()
    ws = wb.active

    # Создаем колонку с заголовками
    cell = ws.cell(row=2, column=1, value="Количество сообщений")
    cell.fill = PatternFill(fgColor=Color(rgb=COLOR_INDEX[5]), fill_type="lightUp")
    cell.alignment = Alignment(horizontal="center")

    cell2 = ws.cell(row=3, column=1, value="Количество открытых тикетов")
    cell2.fill = PatternFill(fgColor=Color(rgb=COLOR_INDEX[5]), fill_type="lightUp")
    cell2.alignment = Alignment(horizontal="center")

    cell3 = ws.cell(row=4, column=1, value="Количество закрытых тикетов")
    cell3.fill = PatternFill(fgColor=Color(rgb=COLOR_INDEX[5]), fill_type="lightUp")
    cell3.alignment = Alignment(horizontal="center")

    ws.column_dimensions[get_column_letter(1)].width = 30

    connect = DatabaseConnection().connect
    cursor = await connect.cursor()

    helpers = await (await cursor.execute("SELECT user_id, full_name FROM helpers")).fetchall()
    month_start_time = (datetime.datetime.now() - datetime.timedelta(days=30)).timestamp()

    # TODO: Можно оптимизировать (таск группа например)
    for helper in range(len(helpers)):
        helper_id = helpers[helper][0]
        full_name = helpers[helper][1]
        column = helper + 2

        ws.column_dimensions[get_column_letter(column)].width = 50
        cell = ws.cell(row=1, column=column, value=full_name)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill(fgColor=Color(rgb=COLOR_INDEX[4]), fill_type="lightUp")

        # Получаем все нужные значения из базы данных
        count_of_messages = (
            await (
                await cursor.execute(
                    "SELECT COUNT(*) FROM messages WHERE helper_id = ? AND time_at > ? ", (
                        helper_id, month_start_time
                    )
                )
            ).fetchone()
        )[0]

        count_of_open_tickets = (
            await (
                await cursor.execute(
                    "SELECT COUNT(*) FROM tickets WHERE start_time > ? AND helper_id = ?", (
                        month_start_time, helper_id
                    )
                )
            ).fetchone()
        )[0]

        count_of_closed_tickets = (
            await (
                await cursor.execute(
                    "SELECT COUNT(*) FROM tickets WHERE end_time > ? AND helper_id = ?", (
                        month_start_time, helper_id
                    )
                )
            ).fetchone()
        )[0]

        cell_messages_count = ws.cell(row=2, column=column, value=str(count_of_messages))
        cell_messages_count.alignment = Alignment(horizontal="center")

        cell_opened_tickets_count = ws.cell(row=3, column=column, value=str(count_of_open_tickets))
        cell_opened_tickets_count.alignment = Alignment(horizontal="center")

        cell_closed_tickets_count = ws.cell(row=4, column=column, value=str(count_of_closed_tickets))
        cell_closed_tickets_count.alignment = Alignment(horizontal="center")

    # Заливаем таблицу в стрим
    file_stream = io.BytesIO()
    wb.save(file_stream)
    wb.close()
    file_stream.seek(0)

    input_file = BufferedInputFile(file_stream.read(), filename="report.xlsx")
    await message.answer_document(input_file, caption=f"Ваша выгрузка за {datetime.datetime.now().month}")

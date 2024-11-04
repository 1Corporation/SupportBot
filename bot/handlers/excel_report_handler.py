import datetime
import os
import io

import openpyxl
from openpyxl.styles import PatternFill, Alignment
from openpyxl.styles.colors import Color, COLOR_INDEX
from openpyxl.utils.cell import get_column_letter
from aiogram.types import Message, BufferedInputFile
from aiogram.filters.command import Command

from bot.dispatcher import dp
from utils import DatabaseConnection


@dp.message(Command("excel_report"))
async def excel_report_handler(message: Message) -> None:
    await send_xlsx_report(message)


async def send_xlsx_report(message: Message) -> None:
    work_day_started_at = int(os.getenv("WORK_DAY_STARTED_AT"))
    work_day_ended_at = int(os.getenv("WORK_DAY_END_AT"))

    wb = openpyxl.Workbook()
    ws = wb.active

    # Сгенерируем колонку с временем
    for i in range(work_day_ended_at - work_day_started_at + 1):
        cell = ws.cell(row=i+2, column=1, value=f"{work_day_started_at + i}:00")
        cell.fill = PatternFill(fgColor=Color(rgb=COLOR_INDEX[5]), fill_type="lightUp")
        cell.alignment = Alignment(horizontal="center")

    connect = DatabaseConnection().connect
    cursor = await connect.cursor()

    # Получаем всех помощников
    all_helpers = await (await cursor.execute("""
    SELECT user_id, full_name FROM helpers
    """)).fetchall()

    for i in range(len(all_helpers)):
        helper = all_helpers[i]
        column = i + 2

        # Подписываем ФИО помощника которое мы достали из базы данных
        cell = ws.cell(row=1, column=column, value=helper[1])
        cell.fill = PatternFill(fgColor=Color(rgb=COLOR_INDEX[4]), fill_type="lightUp")
        cell.alignment = Alignment(horizontal="center")

        # Получаем список сообщений
        list_of_messages = await (await cursor.execute(
            "SELECT time_at FROM messages \
            WHERE helper_id = ? AND time_at > ? \
            ORDER BY time_at", (
                helper[0], work_day_started_at))).fetchall()

        # Проставляем время
        current_hour = work_day_started_at
        for row in list_of_messages:
            hour_of_message = (row[0] - datetime.datetime.now().replace(hour=0, minute=0, second=0).timestamp()) // 3600
            # Если до первого сообщения не было других, проставляем на все часы до этого красный цвет
            if hour_of_message > current_hour:
                while hour_of_message > current_hour:

                    # Даем ячейке значение и стиль
                    cell = ws.cell(
                        row=current_hour - work_day_started_at + 2,
                        column=column,
                        value=f"{current_hour}:00 ❌"
                    )
                    cell.fill = PatternFill(fgColor=Color(rgb=COLOR_INDEX[2]), fill_type="lightUp")
                    cell.alignment = Alignment(horizontal="center")
                    current_hour += 1

            # если в этот час было сообщение, ставим зеленую галочку
            if hour_of_message == current_hour:
                cell = ws.cell(
                    row=current_hour - work_day_started_at + 2,
                    column=column,
                    value=f"{current_hour}:00 ✅"
                )
                cell.fill = PatternFill(fgColor=Color(rgb=COLOR_INDEX[3]), fill_type="lightUp")
                cell.alignment = Alignment(horizontal="center")
                current_hour += 1

        # Если остались часы до конца рабочего дня, добавим красных
        while current_hour < work_day_ended_at + 1:

            # Даем ячейке значение и стиль
            cell = ws.cell(
                row=current_hour - work_day_started_at + 2,
                column=column,
                value=f"{current_hour}:00 ❌"
            )
            cell.fill = PatternFill(fgColor=Color(rgb=COLOR_INDEX[2]), fill_type="lightUp")
            cell.alignment = Alignment(horizontal="center")

            current_hour += 1

        # Делаем колонку больше
        ws.column_dimensions[get_column_letter(i+2)].width = 50

    # Заливаем таблицу в стрим
    file_stream = io.BytesIO()
    wb.save(file_stream)
    wb.close()
    file_stream.seek(0)

    input_file = BufferedInputFile(file_stream.read(), filename="report.xlsx")
    await message.answer_document(input_file, caption=f"Ваша выгрузка за {datetime.datetime.now()}")
